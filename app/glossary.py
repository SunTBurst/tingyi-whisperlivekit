"""Local, deterministic transcript glossary rules and CSV/XLSX interchange."""
from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Any, Iterable


FIELDS = ("enabled", "language", "original", "replacement", "match_mode",
          "target_language", "fixed_translation", "notes")
FIELD_LABELS = {
    "enabled": "启用", "language": "适用语言", "original": "误识别写法",
    "replacement": "标准词", "match_mode": "匹配方式", "target_language": "目标语言",
    "fixed_translation": "固定译名", "notes": "备注",
}
# text: case-insensitive substring; *_case_sensitive: exact case;
# whole_word: English token boundaries (CJK remains literal phrase matching).
MATCH_MODES = ("text", "text_case_sensitive", "whole_word", "whole_word_case_sensitive")
MATCH_MODE_LABELS = {
    "文字匹配（忽略大小写）": "text", "文字匹配（区分大小写）": "text_case_sensitive",
    "完整单词（忽略大小写）": "whole_word",
    "完整单词（区分大小写）": "whole_word_case_sensitive",
}
_LABEL_TO_FIELD = {label: field for field, label in FIELD_LABELS.items()}
_LANGUAGE_ALIASES = {
    "中文": "zh", "汉语": "zh", "chinese": "zh", "普通话": "zh",
    "英语": "en", "英文": "en", "english": "en",
    "阿拉伯语": "ar", "阿拉伯文": "ar", "arabic": "ar",
    "全部": "*", "所有语言": "*", "任意": "*", "all": "*", "any": "*",
    "zho-hans": "zh-hans", "zho_hans": "zh-hans", "zh-hans": "zh-hans",
    "zho-hant": "zh-hant", "zho_hant": "zh-hant", "zh-hant": "zh-hant",
    "eng-latn": "en", "eng_latn": "en", "en-latn": "en",
    "arb-arab": "ar", "arb_arab": "ar", "ar-arab": "ar",
}


def canonical_language(value: Any) -> str:
    """Normalize common UI and NLLB language identifiers without model imports."""
    language = str(value or "*").strip().casefold().replace("_", "-")
    return _LANGUAGE_ALIASES.get(language, language)


def language_matches(scope: str | None, language: str | None) -> bool:
    """Return whether a rule scope applies to an ISO-639 or NLLB language code."""
    scoped = canonical_language(scope)
    actual = canonical_language(language)
    if scoped in ("*", "all", "any"):
        return True
    if scoped == actual:
        return True
    # Plain zh is a family scope and may bind to either simplified or traditional
    # Chinese. A script-specific rule also applies when the source only says zh.
    if scoped == "zh" and actual in ("zh-hans", "zh-hant"):
        return True
    if actual == "zh" and scoped in ("zh-hans", "zh-hant"):
        return True
    return actual.startswith(scoped + "-") or scoped.startswith(actual + "-")


def _compiled_rule_patterns(rule: dict[str, Any]) -> tuple[re.Pattern[str], re.Pattern[str]]:
    escaped = re.escape(rule["original"])
    case_sensitive = rule["match_mode"].endswith("case_sensitive")
    flags = 0 if case_sensitive else re.IGNORECASE
    plain = re.compile(escaped, flags)
    word = re.compile(rf"(?<!\w){escaped}(?!\w)", flags) if rule["match_mode"].startswith("whole_word") else plain
    return plain, word


def _truth(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().casefold() not in ("", "0", "false", "no", "否", "停用")


def normalize_rule(rule: dict[str, Any], row_number: int | None = None) -> dict[str, Any]:
    if not isinstance(rule, dict):
        raise ValueError(f"row {row_number}: rule must be an object")
    translated = {_LABEL_TO_FIELD.get(str(k).strip(), str(k).strip()): v for k, v in rule.items()}
    result = {**translated}
    for field in FIELDS:
        result.setdefault(field, "")
    enabled = result["enabled"]
    result["enabled"] = True if enabled is None or not str(enabled).strip() else _truth(enabled)
    result["language"] = canonical_language(result["language"])
    result["original"] = str(result["original"] or "").strip()
    result["replacement"] = str(result["replacement"] or "").strip()
    result["match_mode"] = MATCH_MODE_LABELS.get(str(result["match_mode"]).strip(),
                                                str(result["match_mode"] or "text").strip())
    for field in ("target_language", "fixed_translation", "notes"):
        result[field] = str(result[field] or "").strip()
    if result["target_language"]:
        result["target_language"] = canonical_language(result["target_language"])
    if not result["original"] or not result["replacement"]:
        raise ValueError(f"row {row_number or '?'}: original and replacement are required")
    if result["match_mode"] not in MATCH_MODES:
        raise ValueError(f"row {row_number or '?'}: unsupported match_mode {result['match_mode']!r}")
    return result


def _key(rule: dict[str, Any]) -> tuple[str, str, str]:
    # A source spelling in one language is one rule identity. Changing its match
    # mode is an overwrite, and conflicting replacements cannot silently coexist.
    return (rule["language"], rule["original"].casefold(), "")


def validate_rules(rules: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    result, seen = [], {}
    for index, raw in enumerate(rules, 1):
        rule = normalize_rule(raw, index)
        key = _key(rule)
        prior = seen.get(key)
        if prior:
            if prior != rule:
                raise ValueError(f"conflict: {rule['original']!r} has different rule values")
            continue
        seen[key] = rule
        result.append(rule)
    return result


def _revision(rules: list[dict[str, Any]], enabled: bool) -> str:
    canonical = json.dumps({"enabled": enabled, "rules": rules}, ensure_ascii=False,
                           sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:16]


class Glossary:
    def __init__(self, rules: list[dict[str, Any]], enabled: bool = True):
        self.rules = validate_rules(rules or [])
        self.enabled = bool(enabled)
        self.revision = _revision(self.rules, self.enabled)
        self._patterns = [_compiled_rule_patterns(rule) for rule in self.rules]

    @staticmethod
    def _applies(rule_language: str, language: str | None) -> bool:
        return language_matches(rule_language, language)

    def correct(self, text: str, language: str | None = None) -> dict[str, Any]:
        original_text = str(text or "")
        candidates = []
        if self.enabled:
            for index, rule in enumerate(self.rules):
                if not rule["enabled"] or not self._applies(rule["language"], language):
                    continue
                mode = rule["match_mode"]
                word_boundary = mode.startswith("whole_word") and canonical_language(language).startswith("en")
                pattern = self._patterns[index][1 if word_boundary else 0]
                for match in pattern.finditer(original_text):
                    candidates.append((match.start(), -(match.end() - match.start()), index,
                                       match.end(), match.group(0), rule))
        candidates.sort(key=lambda candidate: (candidate[0], candidate[1], candidate[2]))
        selected = []
        cursor = -1
        for start, neg_len, index, end, matched, rule in candidates:
            if start < cursor:
                continue
            selected.append((start, end, matched, rule, index))
            cursor = end
        pieces, hits, cursor = [], [], 0
        for start, end, matched, rule, index in selected:
            pieces.extend((original_text[cursor:start], rule["replacement"]))
            hits.append({"rule_index": index, "original": matched,
                         "replacement": rule["replacement"], "start": start, "end": end,
                         "language": rule["language"]})
            cursor = end
        pieces.append(original_text[cursor:])
        return {"original": original_text, "corrected": "".join(pieces), "hits": hits,
                "revision": self.revision}

    def decorate(self, row: dict[str, Any]) -> dict[str, Any]:
        decorated = dict(row)
        source_original = str(row.get("source_original", row.get("text", row.get("source", ""))) or "")
        result = self.correct(source_original, row.get("language") or row.get("source_language"))
        source_corrected = result["corrected"]
        decorated.update({"source_original": source_original, "source_corrected": source_corrected,
                          "source": source_corrected, "glossary_hits": result["hits"],
                          "glossary_revision": self.revision})
        if row.get("translation"):
            bound_revision = row.get("translation_source_revision") or row.get("translation_glossary_revision")
            bound_text = row.get("translation_source_text")
            previous_text = bound_text
            if previous_text is None:
                previous_text = row.get("source_corrected", row.get("source", source_original))
            if bound_text is not None:
                decorated["translation_stale"] = not (
                    bound_revision == self.revision and bound_text == source_corrected)
            elif previous_text != source_corrected:
                # Legacy rows have no explicit translation binding. Their prior
                # displayed/corrected source is the best available evidence.
                decorated["translation_stale"] = True
        return decorated

    def apply_fixed_translations(self, text: str, target_language: str | None) -> str:
        """Replace matching corrected terms in translator output for a target language.

        A fixed term is enforced only when the translated text contains its source
        replacement; this post-processor cannot infer translations that omitted it.
        """
        original_text = str(text or "")
        if not self.enabled:
            return original_text
        candidates = []
        for index, rule in enumerate(self.rules):
            if (not rule["enabled"] or not rule["fixed_translation"] or
                    (rule["target_language"] and not self._applies(rule["target_language"], target_language))):
                continue
            mode = rule["match_mode"]
            escaped = re.escape(rule["replacement"])
            flags = 0 if mode.endswith("case_sensitive") else re.IGNORECASE
            if mode.startswith("whole_word") and canonical_language(target_language).startswith("en"):
                escaped = rf"(?<!\w){escaped}(?!\w)"
            pattern = re.compile(escaped, flags)
            for match in pattern.finditer(original_text):
                candidates.append((match.start(), -(match.end() - match.start()), index,
                                   match.end(), rule["fixed_translation"]))
        candidates.sort(key=lambda candidate: (candidate[0], candidate[1], candidate[2]))
        pieces, cursor = [], 0
        for start, _, _, end, replacement in candidates:
            if start < cursor:
                continue
            pieces.extend((original_text[cursor:start], replacement))
            cursor = end
        pieces.append(original_text[cursor:])
        return "".join(pieces)


def _read_rules(path: str | Path) -> list[dict[str, Any]]:
    path = Path(path)
    if path.suffix.casefold() == ".csv":
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as stream:
                reader = csv.DictReader(stream)
                _validate_headers(reader.fieldnames, path)
                rows = []
                for index, row in enumerate(reader, 2):
                    if not any(str(value or "").strip() for value in row.values()):
                        continue
                    rows.append(normalize_rule(row, index))
        except (OSError, UnicodeError, csv.Error) as exc:
            raise ValueError(f"{path}: {exc}") from exc
        _check_import_conflicts(rows)
        return rows
    if path.suffix.casefold() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX support requires openpyxl; install the approved optional dependency") from exc
        try:
            book = load_workbook(path, read_only=True, data_only=True)
            sheet = book.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            _validate_headers(list(headers) if headers else [], path)
            names = [_LABEL_TO_FIELD.get(str(value).strip(), str(value).strip()) for value in headers]
            rows = [normalize_rule(dict(zip(names, row)), index)
                    for index, row in enumerate(values, 2) if any(value is not None for value in row)]
            book.close()
            _check_import_conflicts(rows)
            return rows
        except (OSError, ValueError) as exc:
            raise ValueError(f"{path}: {exc}") from exc
    raise ValueError(f"{path}: expected a .csv or .xlsx file")


def _validate_headers(headers: list[str] | None, path: Path) -> None:
    names = {_LABEL_TO_FIELD.get(str(name).strip(), str(name).strip()) for name in (headers or [])}
    missing = {"original", "replacement"} - names
    if missing:
        raise ValueError(f"{path}: missing required columns: {', '.join(sorted(missing))}")


def _deduplicate(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output, seen = [], {}
    for row in rows:
        key = _key(row)
        old = seen.get(key)
        if old is not None:
            if old["replacement"] != row["replacement"]:
                raise ValueError(f"conflict in import at row {len(output) + 2}: "
                                 f"{row['original']!r} has different replacements")
            continue
        seen[key] = row
        output.append(row)
    return output


def _check_import_conflicts(rows: list[dict[str, Any]]) -> None:
    seen: dict[tuple[str, str, str], dict[str, Any]] = {}
    for index, row in enumerate(rows, 2):
        key = _key(row)
        old = seen.get(key)
        if old and old != row:
            raise ValueError(f"conflict in import at row {index}: "
                             f"{row['original']!r} has different rule values")
        seen[key] = row


def _write(path: str | Path, rules: list[dict[str, Any]]) -> None:
    path = Path(path)
    normalized = [normalize_rule(rule, index) for index, rule in enumerate(rules, 1)]
    if path.suffix.casefold() == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=[FIELD_LABELS[field] for field in FIELDS])
            writer.writeheader()
            for rule in normalized:
                writer.writerow({FIELD_LABELS[field]: ("是" if rule[field] else "否") if field == "enabled"
                                 else rule[field] for field in FIELDS})
        return
    if path.suffix.casefold() == ".xlsx":
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise RuntimeError("XLSX support requires openpyxl; install the approved optional dependency") from exc
        book = Workbook()
        sheet = book.active
        sheet.title = "Glossary"
        sheet.append([FIELD_LABELS[field] for field in FIELDS])
        for rule in normalized:
            sheet.append([rule[field] for field in FIELDS])
        book.save(path)
        return
    raise ValueError(f"{path}: expected a .csv or .xlsx file")


def export_template(path: str | Path) -> None:
    _write(path, [])


def export_rules(path: str | Path, rules: list[dict[str, Any]]) -> None:
    _write(path, rules)


def preview_import(path: str | Path, current_rules: list[dict[str, Any]]) -> dict[str, Any]:
    incoming = _read_rules(path)
    current = validate_rules(current_rules or [])
    by_key = {_key(rule): rule for rule in current}
    added, conflicts, skipped, unique_incoming = [], [], 0, []
    incoming_seen: dict[tuple[str, str, str], dict[str, Any]] = {}
    for rule in incoming:
        key = _key(rule)
        duplicate = incoming_seen.get(key)
        if duplicate is not None:
            if duplicate != rule:
                raise ValueError(f"conflict in import: {rule['original']!r} has different rule values")
            skipped += 1
            continue
        incoming_seen[key] = rule
        unique_incoming.append(rule)
        old = by_key.get(_key(rule))
        if old is None:
            added.append(rule)
        elif old == rule:
            skipped += 1
        else:
            conflicts.append({"original": rule["original"], "language": rule["language"],
                              "current": old, "incoming": rule})
    return {"added": added, "conflicts": conflicts, "overwritten": conflicts,
            "skipped": skipped, "incoming": unique_incoming,
            "counts": {"added": len(added), "overwritten": len(conflicts), "skipped": skipped}}


def apply_import(current_rules: list[dict[str, Any]], preview: dict[str, Any], mode: str) -> list[dict[str, Any]]:
    if mode not in ("merge", "replace"):
        raise ValueError("mode must be 'merge' or 'replace'")
    if mode == "replace":
        return validate_rules(preview["incoming"])
    result = validate_rules(current_rules or [])
    positions = {_key(rule): index for index, rule in enumerate(result)}
    for rule in preview["incoming"]:
        key = _key(rule)
        if key in positions:
            result[positions[key]] = rule
        else:
            positions[key] = len(result)
            result.append(rule)
    return validate_rules(result)
